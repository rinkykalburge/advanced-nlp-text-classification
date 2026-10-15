"""
Builds the Week 3 report (Word .docx, plus PDF if LibreOffice is installed)
from the results produced by text_classification.py.

Run AFTER text_classification.py:
    pip install python-docx
    python build_report.py            # reads ./outputs, writes ./outputs/Week3_Report.docx

Every number, table and result sentence in the report is read from
outputs/results.json, so the report always matches what your code produced.
"""
import argparse
import json
import os
import shutil
import subprocess

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

AUTHOR = "Sangeeta Kalburge"
COLLEGE = "BKIT Bhalki"


# ----------------------------- helpers ------------------------------------
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(doc, header, rows, widths=None, bold_row=None, font=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "2F4F7F")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font)
            if bold_row is not None and ri == bold_row:
                r.bold = True
                shade(cells[i], "E3EBF6")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def para(doc, text, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    return p


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def figure(doc, path, caption, width=6.0):
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = para(doc, caption, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        c.runs[0].font.size = Pt(9)


def pct(x):
    return f"{100 * x:.2f}%"


# ----------------------------- main ---------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default="outputs")
    a = ap.parse_args()
    D = a.dir
    R = json.load(open(f"{D}/results.json"))
    S, results = R["stats"], R["results"]
    labels = S["labels"]
    by = {r["model"]: r for r in results}
    best = by[R["best_model"]]
    ranked = sorted(results, key=lambda r: r["f1_macro"], reverse=True)
    worst = ranked[-1]
    fastest = min(results, key=lambda r: r["train_seconds"])
    slowest = max(results, key=lambda r: r["train_seconds"])

    # per-class insights for the best model
    pc = best["per_class"]
    best_cls = max(labels, key=lambda l: pc[l]["f1-score"])
    worst_cls = min(labels, key=lambda l: pc[l]["f1-score"])
    cm = best["confusion_matrix"]
    pairs = [(cm[i][j], labels[i], labels[j])
             for i in range(len(labels)) for j in range(len(labels)) if i != j]
    top_pair = max(pairs)
    n_errors = sum(sum(row) for row in cm) - sum(cm[i][i] for i in range(len(cm)))
    gap = best["cv_f1_macro"] - best["f1_macro"]

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # ---------------- title page ----------------
    for _ in range(5):
        doc.add_paragraph()
    para(doc, "NLP Internship - Week 3 Task", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(16)
    t = para(doc, "Implementing and Evaluating Text Classification Models",
             bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    t.runs[0].font.size = Pt(26)
    t.runs[0].font.color.rgb = RGBColor(0x2F, 0x4F, 0x7F)
    para(doc, "Methodology, Findings and Analysis", italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(14)
    for _ in range(4):
        doc.add_paragraph()
    para(doc, f"Prepared by: {AUTHOR}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, COLLEGE, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # ---------------- contents ----------------
    doc.add_heading("Contents", 1)
    for line in ["1. Introduction and Objective", "2. Problem Definition and Data Strategy",
                 "3. Preprocessing and Feature Engineering",
                 "4. Classification Methods and Justification",
                 "5. Experimental Setup", "6. Evaluation Methodology",
                 "7. Results", "8. Analysis and Insights", "9. Challenges Faced",
                 "10. Limitations and Future Improvements", "11. Conclusion",
                 "12. References", "Appendix A - Reproducing the Experiment"]:
        doc.add_paragraph(line)
    doc.add_page_break()

    # ---------------- 1 ----------------
    doc.add_heading("1. Introduction and Objective", 1)
    para(doc, "Text classification is the task of assigning a piece of text to one "
              "of a set of predefined categories. It underpins spam filtering, "
              "news categorisation, support-ticket routing and sentiment analysis. "
              "The objective of this task is to implement several text "
              "classification techniques on publicly available data, evaluate them "
              "with established NLP metrics, and critically analyse the results.")
    para(doc, "This report documents the complete workflow: how the problem was "
              "simulated, how the text was preprocessed and converted to numerical "
              "features, why particular models were chosen, how they were evaluated, "
              "and what was learned. Four classifiers were compared: Multinomial "
              "Naive Bayes (baseline), Logistic Regression, a Linear Support Vector "
              "Machine, and a feed-forward neural network (Multi-Layer Perceptron).")

    # ---------------- 2 ----------------
    doc.add_heading("2. Problem Definition and Data Strategy", 1)
    para(doc, f"Problem: given the body of a message, predict which of "
              f"{len(labels)} topics it belongs to ({', '.join(labels)}). This is a "
              f"single-label, multi-class text classification problem.")
    para(doc, f"Dataset: {S['source']}. The 20 Newsgroups collection is a standard, "
              "publicly available benchmark of Usenet posts distributed with "
              "scikit-learn, which makes the experiment fully reproducible. A "
              "subset of topically distinct and partly overlapping categories was "
              "chosen so that the task is realistic: some classes are easy to "
              "separate while others share vocabulary.")
    para(doc, "Controlled scenario: the headers, signature footers and quoted "
              "replies were removed using the loader's `remove` option. Without this "
              "step the metadata (e-mail addresses, organisation names, "
              "newsgroup-specific signatures) lets models cheat and overestimates "
              "real performance. The standard train/test partition supplied with the "
              "dataset was used, and the test set was kept untouched until the final "
              "evaluation.")
    cc_tr, cc_te = S["train_class_counts"], S["test_class_counts"]
    table(doc, ["Class", "Training documents", "Test documents"],
          [[l, cc_tr.get(l, 0), cc_te.get(l, 0)] for l in labels] +
          [["Total", S["n_train"], S["n_test"]]],
          widths=[2.6, 1.8, 1.6], bold_row=len(labels))
    figure(doc, f"{D}/class_distribution.png", "Figure 1: Class distribution of the training set.", 5.0)

    # ---------------- 3 ----------------
    doc.add_heading("3. Preprocessing and Feature Engineering", 1)
    doc.add_heading("3.1 Text cleaning", 2)
    para(doc, "Each document passes through the following steps before modelling:")
    bullets(doc, [
        "Lower-casing, so that 'Space' and 'space' are the same token.",
        "Removal of e-mail addresses and URLs, which are noise for topic classification.",
        "Removal of digits and punctuation, keeping alphabetic tokens only.",
        "Whitespace normalisation.",
        f"Removal of documents with fewer than 3 words after cleaning "
        f"({S['dropped_train']} training and {S['dropped_test']} test documents were dropped), "
        "because they carry too little signal to classify.",
    ])
    para(doc, f"On average a training document contained {S['avg_words_before']:.0f} "
              f"whitespace-separated tokens before cleaning and {S['avg_words_after']:.0f} after.")
    para(doc, "Example (first 350 characters):", bold=True)
    para(doc, "Before: " + S["example_raw"].replace("\n", " "), italic=True)
    para(doc, "After:  " + S["example_clean"], italic=True)
    doc.add_heading("3.2 Feature engineering: TF-IDF", 2)
    para(doc, "Machine-learning models need numbers, not words. Each cleaned document "
              "was converted into a TF-IDF (Term Frequency - Inverse Document "
              "Frequency) vector. TF measures how often a term occurs in a document; "
              "IDF down-weights terms that appear in many documents (such as 'would' "
              "or 'one') and up-weights discriminative ones (such as 'orbit' or "
              "'goalie'). The settings used were:")
    bullets(doc, [
        "Unigrams and bigrams (ngram_range = (1, 2)) to capture short phrases such as 'space shuttle'.",
        "English stop-word removal to discard function words.",
        "min_df = 2 and max_df = 0.9: ignore terms that occur in only one document (typos, rare names) or in over 90% of documents.",
        "Sublinear TF scaling (1 + log tf), which stops very long documents or repeated words from dominating.",
        "A cap of 20,000 features to limit dimensionality and memory use.",
    ])
    para(doc, f"After fitting on the full training set the vocabulary contained "
              f"{R['n_features']:,} features.")
    doc.add_heading("3.3 Avoiding data leakage", 2)
    para(doc, "The TF-IDF vectoriser is part of the scikit-learn Pipeline, so inside "
              "cross-validation it is fitted only on the training folds of each split "
              "and applied to the held-out fold. Vocabulary and IDF values therefore "
              "never see validation or test text, which keeps the estimates honest.")

    # ---------------- 4 ----------------
    doc.add_heading("4. Classification Methods and Justification", 1)
    table(doc, ["Model", "How it works", "Why it was chosen"], [
        ["Multinomial Naive Bayes (baseline)",
         "Applies Bayes' theorem assuming word occurrences are independent given the class.",
         "Very fast, strong baseline for text, needs little tuning. Shows how much more complex models actually add."],
        ["Logistic Regression",
         "Linear model that passes a weighted sum of TF-IDF features through a softmax to give class probabilities.",
         "Well-calibrated, robust on high-dimensional sparse data, and interpretable: its weights show which words drive each class."],
        ["Linear SVM",
         "Finds the hyperplane that separates classes with the maximum margin.",
         "Widely regarded as one of the best classical methods for sparse, high-dimensional text; text is often close to linearly separable."],
        ["Neural Network (MLP)",
         "Feed-forward network with one hidden layer (ReLU) on top of the TF-IDF vector.",
         "Can learn non-linear feature interactions; included to test whether extra capacity helps on this dataset."],
    ], widths=[1.5, 2.4, 2.6], font=8.5)
    para(doc, "Together these cover a generative model (Naive Bayes), two "
              "discriminative linear models, and a neural approach, so the comparison "
              "shows whether model complexity or feature quality matters more.")

    # ---------------- 5 ----------------
    doc.add_heading("5. Experimental Setup", 1)
    bullets(doc, [
        "Split: the dataset's official train/test partition; the test set is used once, for final reporting.",
        "Hyper-parameter tuning: GridSearchCV with stratified 5-fold cross-validation on the training set only, optimising macro-F1.",
        "Fixed random seed (42) for reproducibility.",
        "Software: Python, scikit-learn, pandas, NumPy and Matplotlib.",
    ])
    rows = []
    for r in results:
        rows.append([r["model"], ", ".join(f"{k} = {v}" for k, v in r["best_params"].items())])
    para(doc, "Best hyper-parameters found by the grid search:", bold=True)
    table(doc, ["Model", "Selected hyper-parameters"], rows, widths=[2.3, 4.2])

    # ---------------- 6 ----------------
    doc.add_heading("6. Evaluation Methodology", 1)
    para(doc, "Because no single number describes a classifier completely, several "
              "complementary metrics were used. For a given class, TP, FP, FN and TN "
              "denote true positives, false positives, false negatives and true negatives.")
    table(doc, ["Metric", "Definition", "What it tells us"], [
        ["Accuracy", "(TP + TN) / all predictions", "Share of documents classified correctly. Can mislead if classes are imbalanced."],
        ["Precision", "TP / (TP + FP)", "Of the documents predicted as a class, how many truly belong to it."],
        ["Recall", "TP / (TP + FN)", "Of the documents that truly belong to a class, how many were found."],
        ["F1-score", "2 x P x R / (P + R)", "Harmonic mean of precision and recall; penalises imbalance between them."],
        ["Macro average", "Unweighted mean over classes", "Treats every class equally; sensitive to weak classes."],
        ["Weighted average", "Mean weighted by class size", "Reflects overall performance on the actual class mix."],
        ["Confusion matrix", "Counts of true vs predicted class", "Shows exactly which classes get confused with each other."],
        ["Cross-validation F1", "Mean macro-F1 over 5 folds", "Estimates how stable the model is on unseen data; compared with the test score to detect over-fitting."],
    ], widths=[1.3, 2.0, 3.2], font=8.5)
    para(doc, "Macro-F1 was chosen as the main model-selection metric because it gives "
              "equal importance to every topic, so a model cannot hide a poor class "
              "behind a large one. Training time was also recorded as a practical cost measure.")

    # ---------------- 7 ----------------
    doc.add_heading("7. Results", 1)
    doc.add_heading("7.1 Overall comparison on the test set", 2)
    rows = [[r["model"], pct(r["accuracy"]), f"{r['precision_macro']:.4f}",
             f"{r['recall_macro']:.4f}", f"{r['f1_macro']:.4f}",
             f"{r['cv_f1_macro']:.4f}", f"{r['train_seconds']:.1f}"] for r in results]
    best_idx = [r["model"] for r in results].index(best["model"])
    table(doc, ["Model", "Accuracy", "Precision (macro)", "Recall (macro)",
                "F1 (macro)", "CV F1 (5-fold)", "Train + tune time (s)"],
          rows, widths=[1.6, 0.8, 0.9, 0.8, 0.8, 0.8, 0.9], bold_row=best_idx, font=8)
    para(doc, "The best model by test macro-F1 is highlighted.", italic=True)
    figure(doc, f"{D}/model_comparison.png", "Figure 2: Accuracy, precision, recall and F1 of all models on the test set.", 6.0)

    doc.add_heading(f"7.2 Per-class results - {best['model']}", 2)
    rows = [[l, f"{pc[l]['precision']:.3f}", f"{pc[l]['recall']:.3f}",
             f"{pc[l]['f1-score']:.3f}", int(pc[l]["support"])] for l in labels]
    table(doc, ["Class", "Precision", "Recall", "F1-score", "Support"], rows,
          widths=[2.4, 1.0, 1.0, 1.0, 1.0])
    figure(doc, f"{D}/confusion_matrix.png", f"Figure 3: Confusion matrix of {best['model']} on the test set.", 4.8)

    doc.add_heading("7.3 Most influential terms", 2)
    para(doc, "The Logistic Regression weights make the learned decision rules visible. "
              "The highest-weighted terms for each class are:")
    for c, words in R["top_features"].items():
        doc.add_paragraph(f"{c}: " + ", ".join(words), style="List Bullet")
    figure(doc, f"{D}/top_features.png", "Figure 4: Highest-weighted terms per class (Logistic Regression).", 6.4)

    # ---------------- 7.4 advanced (optional) ----------------
    adv_path = f"{D}/results_adv.json"
    if os.path.exists(adv_path):
        A = json.load(open(adv_path))
        am = A["models"]
        abest = max(am, key=lambda r: r["f1_macro"])
        doc.add_heading("7.4 Advanced experiments", 2)
        para(doc, f"To go beyond the baseline comparison, a second set of experiments was run on "
                  f"{A['source']} ({A['n_train']} training and {A['n_test']} test documents). It covers "
                  "feature ablation, a stacking ensemble, repeated cross-validation, confidence "
                  "intervals, significance testing and learning curves.")
        doc.add_heading("Feature ablation", 2)
        para(doc, "The same Linear SVM was trained on three feature sets to isolate the effect of "
                  "the representation. The feature set was selected using cross-validation only, never the test set.")
        table(doc, ["Feature set", "# features", "CV macro-F1", "Test accuracy", "Test macro-F1"],
              [[r["features"], f"{r['n_features']:,}", f"{r['cv_f1_macro']:.4f}", pct(r["accuracy"]),
                f"{r['f1_macro']:.4f}"] for r in A["ablation"]],
              widths=[2.4, 0.9, 1.0, 1.1, 1.1], font=8.5)
        para(doc, f"Selected feature set: {A['best_features']}.")
        doc.add_heading("Stronger models and a stacking ensemble", 2)
        para(doc, "Complement Naive Bayes (better suited to skewed text than the multinomial form) was "
                  "added, together with a stacking ensemble in which Logistic Regression, a calibrated "
                  "Linear SVM and Complement NB make predictions and a Logistic Regression meta-model "
                  "learns how to combine them. Stacking can help when base models make different errors.")
        table(doc, ["Model", "Accuracy", "F1 (macro)", "ROC-AUC (macro)", "Repeated CV F1 (5x3)"],
              [[r["model"], pct(r["accuracy"]), f"{r['f1_macro']:.4f}",
                f"{r['roc_auc_macro']:.4f}" if r.get("roc_auc_macro") is not None else "n/a",
                f"{r['rcv_mean']:.4f} +/- {r['rcv_std']:.4f}" if "rcv_mean" in r else "n/a"]
               for r in am], widths=[2.0, 0.9, 0.9, 1.2, 1.5], bold_row=am.index(abest), font=8.5)
        doc.add_heading("Statistical reliability", 2)
        ci = A["bootstrap_ci"]
        para(doc, f"Bootstrap confidence interval. Resampling the test set 1,000 times gives a 95% "
                  f"interval of [{ci['low']:.4f}, {ci['high']:.4f}] for the macro-F1 of {ci['model']}. "
                  "Differences between models smaller than this interval should not be over-interpreted.")
        if A["mcnemar"]:
            para(doc, "McNemar's exact test compares the best model with each other model on the same "
                      "test documents, using only the cases where exactly one of the two is correct "
                      "(b = only the best model is right, c = only the other model is right). A p-value "
                      "below 0.05 indicates a statistically significant difference.")
            table(doc, ["Compared with", "b", "c", "p-value", "Significant at 0.05?"],
                  [[t["other"], t["best_only_correct"], t["other_only_correct"], f"{t['p_value']:.4f}",
                    "Yes" if t["p_value"] < 0.05 else "No"] for t in A["mcnemar"]],
                  widths=[2.4, 0.7, 0.7, 1.0, 1.6], font=8.5)
            nsig = sum(t["p_value"] < 0.05 for t in A["mcnemar"])
            para(doc, f"{nsig} of {len(A['mcnemar'])} comparisons were significant. " +
                 ("Where the test is not significant, the models should be regarded as equally good on this data, "
                  "and the simpler or faster one is preferable." if nsig < len(A["mcnemar"]) else
                  "The best model is significantly better than every alternative."))
        figure(doc, f"{D}/pr_curves.png", "Figure 5: One-vs-rest precision-recall curves of the best advanced model.", 4.6)
        figure(doc, f"{D}/learning_curve.png", "Figure 6: Learning curve of the best advanced model.", 4.6)
        if "learning_curve" in A:
            lc = A["learning_curve"]
            para(doc, f"Learning curve. Validation macro-F1 moved from {lc['val'][0]:.3f} with {lc['sizes'][0]} "
                      f"training documents to {lc['val'][-1]:.3f} with {lc['sizes'][-1]}, while training macro-F1 ended at "
                      f"{lc['train'][-1]:.3f}. " +
                 ("The validation score was still rising at the largest size, so more labelled data would likely help."
                  if lc["val"][-1] - lc["val"][-2] > 0.005 else
                  "The validation score has flattened, so more data of the same kind would add little; better features or models are needed to improve further.") +
                 (" The large train-validation gap indicates some over-fitting." if lc["train"][-1] - lc["val"][-1] > 0.05 else ""))

    # ---------------- 8 ----------------
    doc.add_heading("8. Analysis and Insights", 1)
    para(doc, f"Model ranking. {ranked[0]['model']} achieved the highest test macro-F1 "
              f"({ranked[0]['f1_macro']:.4f}, accuracy {pct(ranked[0]['accuracy'])}), while "
              f"{worst['model']} scored lowest ({worst['f1_macro']:.4f}). The spread between "
              f"best and worst is {ranked[0]['f1_macro'] - worst['f1_macro']:.4f} macro-F1, "
              f"the full order being: " + " > ".join(f"{r['model']} ({r['f1_macro']:.3f})" for r in ranked) + ".")
    para(doc, f"Generalisation. For the best model the cross-validated macro-F1 on training data "
              f"was {best['cv_f1_macro']:.4f} and the test macro-F1 was {best['f1_macro']:.4f}, a "
              f"difference of {gap:+.4f}. " +
         ("A small gap indicates that the tuning procedure produced a model that generalises well "
          "rather than over-fitting the training data."
          if abs(gap) < 0.03 else
          "A gap of this size suggests the test set differs somewhat from the training data "
          "(the official split is chronological) or that some over-fitting to the training folds occurred."))
    para(doc, f"Class-level behaviour. The easiest class for {best['model']} was '{best_cls}' "
              f"(F1 = {pc[best_cls]['f1-score']:.3f}) and the hardest was '{worst_cls}' "
              f"(F1 = {pc[worst_cls]['f1-score']:.3f}). The most frequent confusion was "
              f"'{top_pair[1]}' documents predicted as '{top_pair[2]}' ({top_pair[0]} cases). "
              f"In total {n_errors} of {S['n_test']} test documents were misclassified.")
    para(doc, f"Cost versus accuracy. {fastest['model']} was the cheapest to train and tune "
              f"({fastest['train_seconds']:.1f} s) and {slowest['model']} the most expensive "
              f"({slowest['train_seconds']:.1f} s). Comparing these with the F1 scores shows whether the "
              "extra computation of more complex models was repaid by better predictions.")
    para(doc, "Precision versus recall. Where precision and recall of a class diverge, the model "
              "is either over-predicting the class (high recall, low precision) or being too "
              "conservative (high precision, low recall). The per-class table in Section 7.2 and "
              "the confusion matrix identify which case applies to each topic.")
    if R["errors"]:
        doc.add_heading("Error analysis: sample misclassified documents", 2)
        for e in R["errors"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"True: {e['true']} | Predicted: {e['predicted']}. ").bold = True
            p.add_run(f"\"{e['text']}...\"").italic = True
        para(doc, "Inspecting such documents typically reveals short posts, off-topic "
                  "discussion, or vocabulary shared between categories - cases where even a "
                  "human reader would find the label ambiguous.")
    doc.add_heading("Key insights", 2)
    bullets(doc, [
        ("The classifiers finished within " + f"{ranked[0]['f1_macro'] - worst['f1_macro']:.3f}" +
         " macro-F1 of each other, so the TF-IDF feature representation, rather than the choice of algorithm, did most of the work."
         if ranked[0]['f1_macro'] - worst['f1_macro'] < 0.05 else
         "The classifiers differed noticeably (" + f"{ranked[0]['f1_macro'] - worst['f1_macro']:.3f}" +
         " macro-F1 between best and worst), so the choice of algorithm mattered on top of the TF-IDF features."),
        "Linear models are strong and efficient baselines for sparse text features, and their weights are directly interpretable.",
        "Accuracy alone is not enough: macro-F1, per-class scores and the confusion matrix expose which topics are genuinely hard.",
        "Removing headers, footers and quotes lowers scores compared with the raw dataset but gives a far more honest estimate of real-world performance.",
    ])

    # ---------------- 9 ----------------
    doc.add_heading("9. Challenges Faced", 1)
    bullets(doc, [
        "Metadata leakage: raw newsgroup posts contain headers and signatures that identify the class; these had to be removed so the models learn from content.",
        "High-dimensional sparse features: TF-IDF with bigrams yields tens of thousands of dimensions, requiring a feature cap and sparse-friendly algorithms.",
        "Overlapping vocabulary: some topics share words, producing unavoidable confusions between related classes.",
        "Hyper-parameter selection: choosing settings without touching the test set required nested use of cross-validation inside a Pipeline.",
        "Fair comparison: all models used identical features, splits and tuning metric so that differences reflect the algorithm and not the setup.",
        f"Training cost: {slowest['model']} was the slowest model to tune and train ({slowest['train_seconds']:.1f} s versus {fastest['train_seconds']:.1f} s for {fastest['model']}), which limits how large a hyper-parameter grid is practical.",
    ])

    # ---------------- 10 ----------------
    doc.add_heading("10. Limitations and Future Improvements", 1)
    doc.add_heading("Limitations", 2)
    bullets(doc, [
        f"Only {len(labels)} categories of the full 20 were used, so results will not transfer directly to harder, finer-grained problems.",
        "Bag-of-words features ignore word order and context, so sarcasm, negation and polysemy (words with several meanings) are not captured.",
        "The data are old Usenet posts; modern text (social media, reviews) differs in style and vocabulary.",
        "Results come from a single train/test split and one random seed, so small differences between models may not be statistically significant.",
        "No lemmatisation or stemming was applied.",
    ])
    doc.add_heading("Proposed improvements", 2)
    bullets(doc, [
        "Use pretrained contextual embeddings or fine-tune a transformer such as BERT or DistilBERT.",
        "Add lemmatisation and compare character n-gram features, which handle misspellings.",
        "Repeat experiments over several seeds and report confidence intervals or run a significance test (for example McNemar's test).",
        "Extend to all 20 classes and evaluate on a second, more modern dataset to test robustness.",
        "Handle class imbalance with class weights or resampling if the data are skewed.",
        "Add an error-driven loop: inspect misclassifications, refine preprocessing, and retrain.",
    ])

    # ---------------- 11 ----------------
    doc.add_heading("11. Conclusion", 1)
    para(doc, f"This task built an end-to-end text classification pipeline on {S['source']}: "
              f"cleaning, TF-IDF feature engineering, four classifiers tuned with cross-validation, and "
              f"a multi-metric evaluation. {best['model']} performed best with {pct(best['accuracy'])} "
              f"accuracy and a macro-F1 of {best['f1_macro']:.4f} on unseen test data. The analysis "
              "shows that careful preprocessing and evaluation matter as much as the choice of "
              "algorithm, and it identifies concrete next steps, mainly contextual embeddings and "
              "more rigorous statistical testing, for future iterations.")

    # ---------------- 12 ----------------
    doc.add_heading("12. References", 1)
    for ref in [
        "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "Lang, K. (1995). NewsWeeder: Learning to Filter Netnews. Proceedings of ICML (origin of the 20 Newsgroups dataset).",
        "Manning, C., Raghavan, P., Schutze, H. (2008). Introduction to Information Retrieval. Cambridge University Press.",
        "Joachims, T. (1998). Text Categorization with Support Vector Machines. ECML.",
        "scikit-learn documentation: Working with Text Data and the 20 Newsgroups dataset (scikit-learn.org).",
    ]:
        doc.add_paragraph(ref, style="List Number")

    doc.add_heading("Appendix A - Reproducing the Experiment", 1)
    para(doc, "1. pip install scikit-learn pandas numpy matplotlib python-docx")
    para(doc, "2. python text_classification.py   (trains, tunes, evaluates, saves figures and results.json)")
    para(doc, "3. python text_classification_advanced.py   (optional: advanced experiments -> results_adv.json)")
    para(doc, "4. python build_report.py          (generates this report; includes Section 7.4 if results_adv.json exists)")
    para(doc, "All code is in the accompanying GitHub repository / file text_classification.py.")

    docx_path = f"{D}/Week3_Report.docx"
    doc.save(docx_path)
    print("Saved", docx_path)

    # ---------------- description (200+ words) ----------------
    desc = (
        f"This report documents my Week 3 NLP internship task: implementing and evaluating text "
        f"classification models. I framed a controlled multi-class problem using the public 20 Newsgroups "
        f"dataset, selecting {len(labels)} topics ({', '.join(labels)}), and removed headers, footers and quoted "
        f"text to prevent metadata leakage. The report explains each preprocessing step (lower-casing, "
        f"removing e-mails, URLs, digits and punctuation, filtering very short documents) and the feature "
        f"engineering choice of TF-IDF with unigrams and bigrams, stop-word removal, frequency thresholds and "
        f"sublinear term frequency, producing a vocabulary of {R['n_features']:,} features. "
        f"I implemented four classifiers in a scikit-learn Pipeline: Multinomial Naive Bayes as a baseline, "
        f"Logistic Regression, a Linear SVM and a Multi-Layer Perceptron neural network, and justified why each "
        f"suits sparse, high-dimensional text. Hyper-parameters were tuned with stratified 5-fold "
        f"cross-validation on training data only, optimising macro-F1, and the untouched test set was used once for "
        f"final evaluation. I report accuracy, precision, recall, F1 (macro and weighted), per-class results, "
        f"confusion matrices and training time. The best model was {best['model']} with {pct(best['accuracy'])} "
        f"accuracy and a macro-F1 of {best['f1_macro']:.4f}; the weakest class was '{worst_cls}' and the most common "
        f"confusion was between '{top_pair[1]}' and '{top_pair[2]}'. The analysis covers generalisation "
        f"(cross-validation versus test), the cost versus accuracy trade-off, and error analysis of "
        f"misclassified documents. I also discuss challenges such as data leakage and overlapping vocabulary, "
        f"limitations such as bag-of-words features and a single data split, and future improvements including "
        f"transformer embeddings, lemmatisation and statistical significance testing. The complete, reproducible "
        f"code is provided alongside the report."
    )
    open(f"{D}/description.txt", "w").write(desc)
    print(f"Saved {D}/description.txt ({len(desc.split())} words)")

    # ---------------- PDF (needs LibreOffice) ----------------
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", D, docx_path],
                       check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if os.path.exists(f"{D}/Week3_Report.pdf"):
            print("Saved", f"{D}/Week3_Report.pdf")
    else:
        print("LibreOffice not found: open Week3_Report.docx in Word and use File > Save As > PDF.")


if __name__ == "__main__":
    main()
